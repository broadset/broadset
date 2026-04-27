#!/usr/bin/env python3
"""Generate a python-pptx oracle fixture for Broadset PPTX import tests.

The fixture bytes are generated locally and are not committed. The
companion JSON records geometry, text, and fill expectations from the
source values passed to python-pptx, so the optional Vitest gate can catch
Broadset ECMA-376 parsing regressions against a known-good writer.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
except ImportError as exc:
    raise SystemExit(
        "python-pptx is required. Install it with: "
        "python3 -m pip install -r scripts/requirements-pptx-oracle.txt"
    ) from exc


PACKAGE_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT_DIR = PACKAGE_ROOT / "test-fixtures" / "pptx" / "oracle" / ".cache"
FIXTURE_BASENAME = "python-pptx-oracle"
MM_PER_INCH = 25.4
SLIDE_WIDTH_IN = 10.0
SLIDE_HEIGHT_IN = 7.5


def mm(value_in_inches: float) -> float:
    return round(value_in_inches * MM_PER_INCH, 2)


def rgb(hex_value: str) -> RGBColor:
    normalized = hex_value.removeprefix("#")

    if len(normalized) != 6:
        raise ValueError(f"Expected 6-digit RGB hex, got {hex_value!r}")

    return RGBColor(int(normalized[0:2], 16), int(normalized[2:4], 16), int(normalized[4:6], 16))


def geometry(left: float, top: float, width: float, height: float) -> dict[str, object]:
    return {
        "position": {"x": mm(left), "y": mm(top)},
        "size": {"width": mm(width), "height": mm(height)},
    }


def build_fixture(output_dir: Path) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)

    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_IN)
    presentation.slide_height = Inches(SLIDE_HEIGHT_IN)

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.40), Inches(8.50), Inches(0.80))
    title_box.name = "Oracle Title"
    title_frame = title_box.text_frame
    title_frame.clear()
    title_run = title_frame.paragraphs[0].add_run()
    title_run.text = "python-pptx oracle"
    title_run.font.size = Pt(32)
    title_run.font.bold = True
    title_run.font.name = "Aptos Display"

    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.00), Inches(1.50), Inches(2.50), Inches(1.20))
    rect.name = "Oracle Rectangle"
    rect.fill.solid()
    rect.fill.fore_color.rgb = rgb("#4472C4")
    rect.line.color.rgb = rgb("#1F4E79")

    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4.00), Inches(1.50), Inches(2.00), Inches(1.20))
    oval.name = "Oracle Ellipse"
    oval.fill.solid()
    oval.fill.fore_color.rgb = rgb("#70AD47")
    oval.line.color.rgb = rgb("#548235")

    pptx_path = output_dir / f"{FIXTURE_BASENAME}.pptx"
    expected_path = output_dir / f"{FIXTURE_BASENAME}.expected.json"
    presentation.save(pptx_path)

    expected = {
        "schemaVersion": 1,
        "generator": "python-pptx",
        "fixtureName": FIXTURE_BASENAME,
        "slideCount": 1,
        "canvas": {
            "unit": "mm",
            "width": mm(SLIDE_WIDTH_IN),
            "height": mm(SLIDE_HEIGHT_IN),
        },
        "elements": [
            {
                "type": "text",
                "name": "Oracle Title",
                "textContains": "python-pptx oracle",
                **geometry(0.75, 0.40, 8.50, 0.80),
            },
            {
                "type": "rectangle",
                "name": "Oracle Rectangle",
                "fillHex": "#4472C4",
                **geometry(1.00, 1.50, 2.50, 1.20),
            },
            {
                "type": "ellipse",
                "name": "Oracle Ellipse",
                "fillHex": "#70AD47",
                **geometry(4.00, 1.50, 2.00, 1.20),
            },
        ],
    }

    expected_path.write_text(json.dumps(expected, indent=2, sort_keys=True) + "\n", encoding="utf-8")

    return pptx_path, expected_path


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--out-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help=f"Directory for generated .pptx and .expected.json files (default: {DEFAULT_OUTPUT_DIR})",
    )

    return parser.parse_args(argv)


def main(argv: list[str]) -> int:
    args = parse_args(argv)
    pptx_path, expected_path = build_fixture(args.out_dir)
    print(f"wrote {pptx_path}")
    print(f"wrote {expected_path}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))